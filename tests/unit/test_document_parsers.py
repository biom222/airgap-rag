from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument

from airgap_rag.documents.errors import DocumentParseError
from airgap_rag.documents.parsers import DocxParser, PdfParser, TextParser


def test_text_parser_returns_page_without_number(runtime_path: Path) -> None:
    path = runtime_path / "notes.txt"
    path.write_text("First line\nSecond line", encoding="utf-8")

    parsed = TextParser().parse(path)

    assert len(parsed.pages) == 1
    assert parsed.pages[0].page is None
    assert "Second line" in parsed.pages[0].text


def test_docx_parser_extracts_paragraphs_and_tables(runtime_path: Path) -> None:
    path = runtime_path / "report.docx"
    document = DocxDocument()
    document.add_paragraph("Основной текст")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.save(str(path))

    parsed = DocxParser().parse(path)

    assert parsed.pages[0].page is None
    assert parsed.pages[0].text == "Основной текст\nA\tB"


def test_pdf_parser_preserves_one_based_page_numbers(runtime_path: Path) -> None:
    path = runtime_path / "report.pdf"
    document = fitz.open()
    first = document.new_page()
    first.insert_text((72, 72), "First page")
    second = document.new_page()
    second.insert_text((72, 72), "Second page")
    document.save(path)
    document.close()

    parsed = PdfParser().parse(path)

    assert [page.page for page in parsed.pages] == [1, 2]
    assert "Second page" in parsed.pages[1].text


def test_pdf_parser_wraps_corrupted_pdf_error(runtime_path: Path) -> None:
    path = runtime_path / "broken.pdf"
    path.write_bytes(b"%PDF-not-really-a-pdf")

    with pytest.raises(DocumentParseError, match="PDF"):
        PdfParser().parse(path)
