from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal, Protocol

import fitz
from docx import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from docx.table import Table
from docx.text.paragraph import Paragraph

from airgap_rag.documents.errors import DocumentParseError
from airgap_rag.documents.validation import DOCX_MIME, PDF_MIME, TEXT_MIME


@dataclass(frozen=True, slots=True)
class ParsedBlock:
    kind: Literal["heading", "paragraph", "list", "table", "image"]
    text: str
    section: str | None = None


@dataclass(frozen=True, slots=True)
class ParsedPage:
    page: int | None
    text: str
    blocks: tuple[ParsedBlock, ...] = ()


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    pages: tuple[ParsedPage, ...]


class DocumentParser(Protocol):
    mime_type: str

    def parse(self, path: Path) -> ParsedDocument: ...


class PdfParser:
    mime_type = PDF_MIME

    def __init__(
        self,
        *,
        ocr_enabled: bool = True,
        ocr_languages: str = "rus+eng",
        ocr_min_text_chars: int = 40,
        ocr_dpi: int = 300,
    ) -> None:
        self._ocr_enabled = ocr_enabled
        self._ocr_languages = ocr_languages
        self._ocr_min_text_chars = ocr_min_text_chars
        self._ocr_dpi = ocr_dpi

    def parse(self, path: Path) -> ParsedDocument:
        try:
            with fitz.open(path) as document:
                if document.needs_pass:
                    raise DocumentParseError("Зашифрованные PDF не поддерживаются.")
                pages = tuple(
                    self._parse_page(page, index + 1) for index, page in enumerate(document)
                )
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError("Не удалось разобрать PDF.") from exc  # noqa: RUF001
        return ParsedDocument(pages=pages)

    def _parse_page(self, page: fitz.Page, page_number: int) -> ParsedPage:
        text = page.get_text("text", sort=True)
        if self._ocr_enabled and len(text.strip()) < self._ocr_min_text_chars:
            try:
                text_page = page.get_textpage_ocr(
                    language=self._ocr_languages,
                    dpi=self._ocr_dpi,
                    full=True,
                )
                text = page.get_text("text", textpage=text_page, sort=True)
            except RuntimeError as exc:
                if not text.strip():
                    raise DocumentParseError(
                        "OCR не выполнен: проверьте установку Tesseract и языковых пакетов."
                    ) from exc

        blocks = self._pdf_blocks(page)
        if not blocks and text.strip():
            blocks = (ParsedBlock(kind="paragraph", text=text.strip()),)
        return ParsedPage(page=page_number, text=text, blocks=blocks)

    @staticmethod
    def _pdf_blocks(page: fitz.Page) -> tuple[ParsedBlock, ...]:
        raw: dict[str, Any] = page.get_text("dict", sort=True)
        spans: list[tuple[float, str, float, int]] = []
        for block in raw.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                line_text = "".join(span.get("text", "") for span in line.get("spans", [])).strip()
                if not line_text:
                    continue
                line_spans = line.get("spans", [])
                size = max((float(span.get("size", 0)) for span in line_spans), default=0.0)
                flags = max((int(span.get("flags", 0)) for span in line_spans), default=0)
                spans.append((float(line.get("bbox", (0, 0, 0, 0))[1]), line_text, size, flags))
        if not spans:
            return ()
        body_sizes = sorted(item[2] for item in spans if item[2] > 0)
        median_size = body_sizes[len(body_sizes) // 2] if body_sizes else 0.0
        section: str | None = None
        blocks: list[ParsedBlock] = []
        for _, line_text, size, flags in sorted(spans, key=lambda item: item[0]):
            is_heading = size > median_size * 1.18 or bool(flags & 16)
            if is_heading and len(line_text) <= 255:
                section = line_text
                blocks.append(ParsedBlock(kind="heading", text=line_text, section=section))
            else:
                blocks.append(ParsedBlock(kind="paragraph", text=line_text, section=section))
        try:
            tables = page.find_tables()
            for table in tables.tables:
                rows = [tuple(str(cell or "").strip() for cell in row) for row in table.extract()]
                rows = [row for row in rows if any(row)]
                if rows:
                    markdown = "\n".join("| " + " | ".join(row) + " |" for row in rows)
                    blocks.append(ParsedBlock(kind="table", text=markdown, section=section))
        except (AttributeError, ValueError):
            pass
        for image_index, _ in enumerate(page.get_images(full=True), start=1):
            blocks.append(
                ParsedBlock(
                    kind="image",
                    text=f"[Изображение {image_index} на странице]",
                    section=section,
                )
            )
        return tuple(blocks)


class DocxParser:
    mime_type = DOCX_MIME

    def parse(self, path: Path) -> ParsedDocument:
        try:
            document = DocxDocument(str(path))
            text_blocks: list[str] = []
            parsed_blocks: list[ParsedBlock] = []
            section: str | None = None
            for child in document.element.body.iterchildren():
                if child.tag.endswith("}p"):
                    paragraph = Paragraph(child, document)
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    style = paragraph.style.name if paragraph.style is not None else ""
                    if style.lower().startswith(("heading", "заголовок")):
                        section = text
                        kind: Literal["heading", "paragraph", "list"] = "heading"
                    elif style.lower().startswith(("list", "список")):
                        kind = "list"
                    else:
                        kind = "paragraph"
                    text_blocks.append(text)
                    parsed_blocks.append(ParsedBlock(kind=kind, text=text, section=section))
                elif child.tag.endswith("}tbl"):
                    table = Table(child, document)
                    rows = [tuple(cell.text.strip() for cell in row.cells) for row in table.rows]
                    rows = [row for row in rows if any(row)]
                    if rows:
                        text_blocks.extend("\t".join(row) for row in rows)
                        markdown = "\n".join("| " + " | ".join(row) + " |" for row in rows)
                        parsed_blocks.append(
                            ParsedBlock(kind="table", text=markdown, section=section)
                        )
            for relationship in document.part.rels.values():
                if relationship.reltype == RELATIONSHIP_TYPE.IMAGE:
                    name = Path(relationship.target_ref).name
                    placeholder = f"[Изображение: {name}]"
                    text_blocks.append(placeholder)
                    parsed_blocks.append(
                        ParsedBlock(kind="image", text=placeholder, section=section)
                    )
                elif relationship.reltype.endswith("/footnotes"):
                    root = parse_xml(relationship.target_part.blob)
                    for footnote in root.xpath("//*[local-name()='footnote']"):
                        text = " ".join(
                            part.strip() for part in footnote.itertext() if part.strip()
                        )
                        if text:
                            rendered = f"Сноска: {text}"
                            text_blocks.append(rendered)
                            parsed_blocks.append(
                                ParsedBlock(kind="paragraph", text=rendered, section=section)
                            )
        except Exception as exc:
            raise DocumentParseError("Не удалось разобрать DOCX.") from exc  # noqa: RUF001
        return ParsedDocument(
            pages=(
                ParsedPage(
                    page=None,
                    text="\n".join(text_blocks),
                    blocks=tuple(parsed_blocks),
                ),
            )
        )


class TextParser:
    mime_type = TEXT_MIME

    def parse(self, path: Path) -> ParsedDocument:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except (OSError, UnicodeDecodeError) as exc:
            raise DocumentParseError("Не удалось разобрать TXT.") from exc  # noqa: RUF001
        paragraphs = tuple(
            ParsedBlock(kind="paragraph", text=part.strip())
            for part in text.split("\n\n")
            if part.strip()
        )
        return ParsedDocument(pages=(ParsedPage(page=None, text=text, blocks=paragraphs),))


class ParserRegistry:
    def __init__(
        self,
        parsers: tuple[DocumentParser, ...] | None = None,
        *,
        pdf_parser: PdfParser | None = None,
    ) -> None:
        parser_items = parsers or (pdf_parser or PdfParser(), DocxParser(), TextParser())
        self._parsers = {parser.mime_type: parser for parser in parser_items}

    def get(self, mime_type: str) -> DocumentParser:
        try:
            return self._parsers[mime_type]
        except KeyError as exc:
            raise DocumentParseError(f"Нет parser для MIME-типа {mime_type}.") from exc
